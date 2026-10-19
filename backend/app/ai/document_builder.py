import sys
import os
from pathlib import Path
from typing import Dict, Any, Tuple

BASE_DIR = Path(__file__).resolve().parent.parent.parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from backend.app.config import settings

class DocumentBuilder:
    @staticmethod
    def build_pdf(resume_data: Dict[str, Any], output_path: str) -> str:
        """Generates an ATS-compliant, beautifully styled PDF resume."""
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        doc = SimpleDocTemplate(
            output_path,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        
        # Custom Typography Styles
        name_style = ParagraphStyle(
            'ResumeName',
            fontName='Helvetica-Bold',
            fontSize=22,
            leading=26,
            alignment=1, # Center
            textColor=colors.HexColor("#1e293b")
        )
        
        headline_style = ParagraphStyle(
            'ResumeHeadline',
            fontName='Helvetica',
            fontSize=11,
            leading=14,
            alignment=1,
            textColor=colors.HexColor("#475569")
        )
        
        contact_style = ParagraphStyle(
            'ResumeContact',
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            alignment=1,
            textColor=colors.HexColor("#64748b")
        )

        section_heading_style = ParagraphStyle(
            'SectionHeading',
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#0f172a"),
            spaceBefore=8,
            spaceAfter=2
        )

        body_style = ParagraphStyle(
            'ResumeBody',
            fontName='Helvetica',
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#334155")
        )

        bold_item_style = ParagraphStyle(
            'BoldItem',
            fontName='Helvetica-Bold',
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#1e293b")
        )

        date_style = ParagraphStyle(
            'DateItem',
            fontName='Helvetica-Oblique',
            fontSize=9,
            leading=13,
            alignment=2, # Right align
            textColor=colors.HexColor("#64748b")
        )

        bullet_style = ParagraphStyle(
            'BulletText',
            fontName='Helvetica',
            fontSize=9,
            leading=12.5,
            leftIndent=12,
            firstLineIndent=-8,
            textColor=colors.HexColor("#334155"),
            spaceAfter=2
        )

        story = []

        # Header
        name = resume_data.get("name", "Ayush Rao")
        headline = resume_data.get("headline", "Software Engineer – AI/ML & Data Engineering")
        story.append(Paragraph(name, name_style))
        story.append(Paragraph(headline, headline_style))
        story.append(Spacer(1, 3))

        contact = resume_data.get("contact", {})
        contact_parts = []
        if contact.get("phone"): contact_parts.append(contact["phone"])
        if contact.get("email"): contact_parts.append(contact["email"])
        if contact.get("location"): contact_parts.append(contact["location"])
        if contact.get("linkedin"): contact_parts.append("linkedin.com/in/ayush-ar-rao")
        if contact.get("github"): contact_parts.append("github.com/Ayush2412Rao")
        story.append(Paragraph(" &nbsp;|&nbsp; ".join(contact_parts), contact_style))
        story.append(Spacer(1, 8))

        def add_section_divider(title):
            story.append(Paragraph(title.upper(), section_heading_style))
            story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor("#cbd5e1"), spaceBefore=1, spaceAfter=5))

        # Professional Summary
        summary = resume_data.get("summary", "")
        if summary:
            add_section_divider("Professional Summary")
            story.append(Paragraph(summary, body_style))
            story.append(Spacer(1, 4))

        # Experience
        experience = resume_data.get("experience", [])
        if experience:
            add_section_divider("Experience")
            for exp in experience:
                comp = exp.get("company", "")
                role = exp.get("role", "")
                dates = f"{exp.get('start_date', '')} – {exp.get('end_date', '')}"
                loc = exp.get("location", "")

                left_text = f"<b>{comp}</b> &nbsp;|&nbsp; <i>{role}</i>"
                right_text = f"{dates} &nbsp;|&nbsp; {loc}" if loc else dates

                header_table = Table([[
                    Paragraph(left_text, bold_item_style),
                    Paragraph(right_text, date_style)
                ]], colWidths=[360, 180])
                header_table.setStyle(TableStyle([
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('LEFTPADDING', (0,0), (-1,-1), 0),
                    ('RIGHTPADDING', (0,0), (-1,-1), 0),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 1),
                ]))
                story.append(header_table)

                for b in exp.get("bullets", []):
                    story.append(Paragraph(f"• &nbsp; {b}", bullet_style))
                story.append(Spacer(1, 4))

        # Projects
        projects = resume_data.get("projects", [])
        if projects:
            add_section_divider("Projects")
            for proj in projects:
                title = proj.get("title", "")
                techs = ", ".join(proj.get("technologies", []))
                dates = proj.get("dates", "")

                left_text = f"<b>{title}</b> &nbsp;| &nbsp;<i>{techs}</i>"
                header_table = Table([[
                    Paragraph(left_text, bold_item_style),
                    Paragraph(dates, date_style)
                ]], colWidths=[380, 160])
                header_table.setStyle(TableStyle([
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('LEFTPADDING', (0,0), (-1,-1), 0),
                    ('RIGHTPADDING', (0,0), (-1,-1), 0),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 1),
                ]))
                story.append(header_table)

                for b in proj.get("bullets", []):
                    story.append(Paragraph(f"• &nbsp; {b}", bullet_style))
                story.append(Spacer(1, 4))

        # Education
        education = resume_data.get("education", [])
        if education:
            add_section_divider("Education")
            for edu in education:
                inst = edu.get("institution", "")
                degree = edu.get("degree", "")
                dates = edu.get("dates", "")
                cgpa = f" | CGPA: {edu.get('cgpa')}" if edu.get("cgpa") else ""
                loc = f" | {edu.get('location')}" if edu.get("location") else ""

                left_text = f"<b>{inst}</b> — {degree}{cgpa}"
                right_text = f"{dates}{loc}"

                header_table = Table([[
                    Paragraph(left_text, bold_item_style),
                    Paragraph(right_text, date_style)
                ]], colWidths=[380, 160])
                header_table.setStyle(TableStyle([
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('LEFTPADDING', (0,0), (-1,-1), 0),
                    ('RIGHTPADDING', (0,0), (-1,-1), 0),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 1),
                ]))
                story.append(header_table)
            story.append(Spacer(1, 4))

        # Technical Skills
        skills = resume_data.get("skills", {})
        if skills:
            add_section_divider("Technical Skills")
            if isinstance(skills, dict):
                for cat, skill_items in skills.items():
                    item_str = ", ".join(skill_items) if isinstance(skill_items, list) else str(skill_items)
                    story.append(Paragraph(f"<b>{cat}:</b> &nbsp; {item_str}", body_style))
            elif isinstance(skills, list):
                story.append(Paragraph(", ".join(skills), body_style))
            story.append(Spacer(1, 4))

        # Certifications
        certifications = resume_data.get("certifications", [])
        if certifications:
            add_section_divider("Certifications")
            cert_str = " &nbsp;|&nbsp; ".join(certifications)
            story.append(Paragraph(cert_str, body_style))

        doc.build(story)
        return output_path

    @staticmethod
    def build_docx(resume_data: Dict[str, Any], output_path: str) -> str:
        """Generates an ATS-compliant editable DOCX resume."""
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        doc = Document()

        # Margins: 0.5 inch
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)

        # Name & Title
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_p.add_run(resume_data.get("name", "Ayush Rao"))
        run.bold = True
        run.font.size = Pt(20)
        run.font.name = 'Calibri'

        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub_p.add_run(resume_data.get("headline", ""))
        sub_run.font.size = Pt(11)
        sub_run.font.name = 'Calibri'

        # Contact Info
        contact = resume_data.get("contact", {})
        contact_line = " | ".join(filter(None, [
            contact.get("phone"),
            contact.get("email"),
            contact.get("location"),
            contact.get("linkedin"),
            contact.get("github")
        ]))
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c_run = cp.add_run(contact_line)
        c_run.font.size = Pt(9.5)

        def add_heading(text):
            h = doc.add_paragraph()
            hrun = h.add_run(text.upper())
            hrun.bold = True
            hrun.font.size = Pt(11)
            hrun.font.color.rgb = RGBColor(15, 23, 42)

        # Summary
        if resume_data.get("summary"):
            add_heading("Professional Summary")
            p = doc.add_paragraph(resume_data.get("summary"))
            p.runs[0].font.size = Pt(10)

        # Experience
        if resume_data.get("experience"):
            add_heading("Experience")
            for exp in resume_data.get("experience", []):
                p = doc.add_paragraph()
                r1 = p.add_run(f"{exp.get('company')} – {exp.get('role')} ({exp.get('start_date')} – {exp.get('end_date')})")
                r1.bold = True
                r1.font.size = Pt(10)
                for b in exp.get("bullets", []):
                    bp = doc.add_paragraph(style='List Bullet')
                    brun = bp.add_run(b)
                    brun.font.size = Pt(9.5)

        # Projects
        if resume_data.get("projects"):
            add_heading("Projects")
            for proj in resume_data.get("projects", []):
                p = doc.add_paragraph()
                r1 = p.add_run(f"{proj.get('title')} | {', '.join(proj.get('technologies', []))}")
                r1.bold = True
                r1.font.size = Pt(10)
                for b in proj.get("bullets", []):
                    bp = doc.add_paragraph(style='List Bullet')
                    brun = bp.add_run(b)
                    brun.font.size = Pt(9.5)

        # Education
        if resume_data.get("education"):
            add_heading("Education")
            for edu in resume_data.get("education", []):
                p = doc.add_paragraph()
                cgpa = f" | CGPA: {edu.get('cgpa')}" if edu.get("cgpa") else ""
                r1 = p.add_run(f"{edu.get('institution')} — {edu.get('degree')}{cgpa} ({edu.get('dates')})")
                r1.bold = True
                r1.font.size = Pt(10)

        # Skills
        if resume_data.get("skills"):
            add_heading("Technical Skills")
            skills = resume_data.get("skills", {})
            if isinstance(skills, dict):
                for cat, items in skills.items():
                    p = doc.add_paragraph()
                    r1 = p.add_run(f"{cat}: ")
                    r1.bold = True
                    r1.font.size = Pt(9.5)
                    r2 = p.add_run(", ".join(items) if isinstance(items, list) else str(items))
                    r2.font.size = Pt(9.5)

        # Certifications
        if resume_data.get("certifications"):
            add_heading("Certifications")
            p = doc.add_paragraph(" | ".join(resume_data.get("certifications", [])))
            p.runs[0].font.size = Pt(9.5)

        doc.save(output_path)
        return output_path

document_builder = DocumentBuilder()

if __name__ == "__main__":
    import json
    with open(BASE_DIR / "data" / "master_profile.json", "r", encoding="utf-8") as f:
        data = json.load(f)
    pdf_out = str(BASE_DIR / "data" / "resumes" / "Ayush_Rao_Master_Resume.pdf")
    docx_out = str(BASE_DIR / "data" / "resumes" / "Ayush_Rao_Master_Resume.docx")
    document_builder.build_pdf(data, pdf_out)
    document_builder.build_docx(data, docx_out)
    print(f"Generated PDF: {pdf_out}")
    print(f"Generated DOCX: {docx_out}")
